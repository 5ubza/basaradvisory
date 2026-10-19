"""Rebuild the Elm 591284 application pack from tailored.json + the master CV (profile/cv.yaml).
Run from the jobhunt/ directory:  python applications/2026-10-19_elm_consultant_591284/build.py
Tailoring was authored in-session (no ANTHROPIC_API_KEY) and is passed through the same
validate_against_master guard the pipeline uses, so every bullet traces to a master-CV bullet id."""
from __future__ import annotations

import json
import sys
from datetime import date
from pathlib import Path

import yaml

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
sys.path.insert(0, str(ROOT))

from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import mm  # noqa: E402
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer  # noqa: E402

from jobhunt.config import load_config  # noqa: E402
from jobhunt.models import Job  # noqa: E402
from jobhunt.render import render_docx, render_notes_md, render_pdf  # noqa: E402
from jobhunt.scoring import score_all  # noqa: E402
from jobhunt.tailor import TailoredCV, validate_against_master  # noqa: E402

BASE = "Elm_Consultant-Corporate-Performance-Mgmt_Ibrahim_Alsalem"
JOB_URL = "https://careers.elm.sa/ (job code 591284, posted 18 Oct 2026)"


def cover_letter_paragraphs(cv: dict) -> tuple[list[str], list[str], str]:
    c = cv["contact"]
    header = [cv["name"], f"{c['location']} · {c['phone']} · {c['email']} · {c['linkedin']}",
              date(2026, 10, 19).strftime("%d %B %Y"), "Talent Acquisition, Elm Company, Riyadh",
              "Re: Consultant, Corporate Performance Management (Performance Management Unit), Job Code 591284"]
    body = [
        "Dear Hiring Team,",
        "I am applying for the Consultant role in Elm's Corporate Performance Management department. I know Elm from "
        "the inside: as Senior Analyst, Programs & Planning (2019-2020) I designed KPI scorecards and consolidated "
        "enterprise-wide budget forecasts, financial performance data, and program delivery timelines into integrated "
        "reporting for senior leadership. That is the same work this role is built around: translating objectives into "
        "measurable indicators, validating the numbers behind them, and presenting the result so that decisions can be "
        "made on it.",
        "Since leaving Elm I have done that work at national scale. As Economic Expert at the Ministry of Economy and "
        "Planning I was lead researcher for the labor market team, owning three flagship analytical workstreams and "
        "integrating multi-source administrative data collected in coordination with more than ten government bodies. "
        "The collection, clarification, and follow-up with data owners that this role requires was a daily part of that "
        "job, and the findings fed into published national strategy documents and decision-ready reports for "
        "ministerial leadership and inter-ministerial committees. At the Saudi Tourism Authority I built executive "
        "dashboards and economic impact analyses supporting multi-million SAR investment decisions, and at Strategic "
        "Gears I delivered organizational performance reviews for public and private-sector clients.",
        "What I would bring to the Performance Management Unit is an evidence-first habit: analyzing KPI results and "
        "trends with proper quantitative methods, isolating variances and emerging gaps rather than reporting them, and "
        "turning the analysis into clear tables, visuals, and executive summaries in Power BI, Tableau, and Excel. As a "
        "native bilingual Arabic and English speaker I am comfortable coordinating with business units across Elm and "
        "tracking commitments through quarterly and annual review cycles.",
        "Two points of candor. I do not yet hold a Balanced Scorecard certification; my scorecard experience is hands-on "
        "rather than formally credentialed, and I would welcome formalizing it within Elm's BSC cycle. My experience "
        "also exceeds the four-to-six-year band in the posting. I am applying because the role's remit, corporate "
        "performance measurement at a company that has become the Kingdom's digital enabler, is precisely the work I "
        "want to do, and I am comfortable with the grade.",
        "I am based in Riyadh and available for interview at your convenience. My CV is attached, and I would welcome "
        "the chance to discuss how this experience maps to the unit's annual scorecard and reporting cycle.",
    ]
    signoff = "Kind regards,\nIbrahim H. Alsalem, ChE"
    return header, body, signoff


def write_cover_letter(cv: dict, out: Path) -> list[Path]:
    header, body, signoff = cover_letter_paragraphs(cv)
    # Markdown
    md = out / f"{BASE}_Cover_Letter.md"
    md.write_text("\n\n".join(["  \n".join(header)] + body + [signoff.replace("\n", "  \n")]) + "\n", "utf-8")
    # DOCX
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(48)
        s.left_margin = s.right_margin = Pt(56)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    p = doc.add_paragraph()
    p.add_run(header[0]).bold = True
    for line in header[1:]:
        doc.add_paragraph(line)
    for para in body:
        doc.add_paragraph(para)
    for line in signoff.split("\n"):
        doc.add_paragraph(line)
    docx_path = out / f"{BASE}_Cover_Letter.docx"
    doc.save(str(docx_path))
    # PDF
    ss = getSampleStyleSheet()
    st = ParagraphStyle("b", parent=ss["Normal"], fontName="Helvetica", fontSize=10.5, leading=14, spaceAfter=8)
    bold = ParagraphStyle("h", parent=st, fontName="Helvetica-Bold", spaceAfter=2)
    small = ParagraphStyle("s", parent=st, fontSize=9.5, spaceAfter=2)
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731
    story = [Paragraph(esc(header[0]), bold)] + [Paragraph(esc(h), small) for h in header[1:]] + [Spacer(1, 10)]
    story += [Paragraph(esc(b), st) for b in body]
    story += [Paragraph(esc(line), small) for line in signoff.split("\n")]
    pdf_path = out / f"{BASE}_Cover_Letter.pdf"
    SimpleDocTemplate(str(pdf_path), pagesize=A4, leftMargin=20 * mm, rightMargin=20 * mm, topMargin=18 * mm,
                      bottomMargin=18 * mm, title="Cover letter - Elm 591284", author=cv["name"]).build(story)
    return [md, docx_path, pdf_path]


def main() -> None:
    cfg = load_config(str(ROOT / "config.yaml"))
    cv = yaml.safe_load((ROOT / cfg["paths"]["cv"]).read_text("utf-8")) if not Path(cfg["paths"]["cv"]).is_absolute() \
        else yaml.safe_load(Path(cfg["paths"]["cv"]).read_text("utf-8"))
    raw = json.loads((HERE / "tailored.json").read_text("utf-8"))
    t = validate_against_master(TailoredCV.model_validate(raw), cv)
    kept = {b.source_id for r in t.roles for b in r.bullets}
    asked = {b["source_id"] for r in raw["roles"] for b in r["bullets"]}
    if asked - kept:
        raise SystemExit(f"validator dropped bullets: {sorted(asked - kept)}")
    job = Job(source="manual", title="Consultant - Corporate Performance Management", company="Elm",
              location="Riyadh, SA", url=JOB_URL, posted="2026-10-18",
              description=(HERE / "posting.txt").read_text("utf-8"))
    score_all([job], cfg)
    files = [render_docx(cv, t, HERE / f"{BASE}_CV.docx"), render_pdf(cv, t, HERE / f"{BASE}_CV.pdf"),
             render_notes_md(job, t, "in-session (validated)", HERE / f"{BASE}_notes.md")]
    files += write_cover_letter(cv, HERE)
    for f in files:
        print(f.relative_to(ROOT))


if __name__ == "__main__":
    main()
