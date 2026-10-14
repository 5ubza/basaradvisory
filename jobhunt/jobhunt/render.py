"""Render a TailoredCV to DOCX (python-docx), PDF (reportlab) and Markdown. ATS-friendly: single column, real
headings, no tables/text boxes/images."""
from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

from .tailor import TailoredCV


def slugify(s: str, n: int = 40) -> str:
    return re.sub(r"[^A-Za-z0-9]+", "-", s).strip("-")[:n] or "role"


def _contact_line(cv: dict) -> str:
    c = cv["contact"]
    return " · ".join([c["location"], c["phone"], c["email"], c["linkedin"], c["website"]])


def _role_header(role: dict) -> tuple[str, str]:
    left = f"{role['title']} — {role['company']}" + (f" · {role['company_note']}" if role.get("company_note") else "")
    right = f"{role['dates']} · {role['location']}"
    return left, right


def _sections(cv: dict, t: TailoredCV):
    roles = {r["id"]: r for r in cv["experience"]}
    comps = {g["group"]: g["items"] for g in cv["competencies"]}
    ordered_comps = [(g, comps[g]) for g in t.competency_order if g in comps]
    return roles, ordered_comps


# ------------------------------------------------------------------------------------------ DOCX
def render_docx(cv: dict, t: TailoredCV, path: Path) -> Path:
    roles, comps = _sections(cv, t)
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Pt(40)
        s.left_margin = s.right_margin = Pt(48)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(cv["name"].upper())
    r.bold = True
    r.font.size = Pt(16)
    p = doc.add_paragraph(t.headline)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(_contact_line(cv))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(9)

    doc.add_heading("Professional Summary", level=2)
    doc.add_paragraph(t.summary)

    doc.add_heading("Core Competencies", level=2)
    if t.highlighted_skills:
        p = doc.add_paragraph()
        p.add_run("Key match: ").bold = True
        p.add_run(" · ".join(t.highlighted_skills))
    for g, items in comps:
        p = doc.add_paragraph()
        p.add_run(f"{g}: ").bold = True
        p.add_run(" · ".join(items))

    doc.add_heading("Professional Experience", level=2)
    for tr in t.roles:
        role = roles[tr.role_id]
        left, right = _role_header(role)
        p = doc.add_paragraph()
        p.add_run(left).bold = True
        p.add_run(f"    {right}").italic = True
        for b in tr.bullets:
            doc.add_paragraph(b.text, style="List Bullet")

    doc.add_heading("Education", level=2)
    for e in cv["education"]:
        p = doc.add_paragraph()
        p.add_run(f"{e['degree']} — {e['institution']}").bold = True
        p.add_run(f"    {e['dates']} · {e['location']}").italic = True
        if e.get("note"):
            doc.add_paragraph(e["note"], style="List Bullet")

    doc.add_heading("Credentials & Languages", level=2)
    for c in cv["credentials"]:
        doc.add_paragraph(c, style="List Bullet")
    doc.add_paragraph(f"Languages: {cv['languages']}", style="List Bullet")
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))
    return path


# ------------------------------------------------------------------------------------------ PDF
def render_pdf(cv: dict, t: TailoredCV, path: Path) -> Path:
    roles, comps = _sections(cv, t)
    ss = getSampleStyleSheet()
    body = ParagraphStyle("body", parent=ss["Normal"], fontName="Helvetica", fontSize=9.5, leading=12.5)
    h = ParagraphStyle("h", parent=body, fontName="Helvetica-Bold", fontSize=11, spaceBefore=8, spaceAfter=3)
    name = ParagraphStyle("name", parent=body, fontName="Helvetica-Bold", fontSize=15, alignment=TA_CENTER, leading=18)
    center = ParagraphStyle("c", parent=body, alignment=TA_CENTER, fontSize=8.8)
    esc = lambda s: s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")  # noqa: E731

    story = [Paragraph(esc(cv["name"].upper()), name), Paragraph(esc(t.headline), center),
             Paragraph(esc(_contact_line(cv)), center), Spacer(1, 4),
             Paragraph("PROFESSIONAL SUMMARY", h), Paragraph(esc(t.summary), body),
             Paragraph("CORE COMPETENCIES", h)]
    if t.highlighted_skills:
        story.append(Paragraph("<b>Key match:</b> " + esc(" · ".join(t.highlighted_skills)), body))
    for g, items in comps:
        story.append(Paragraph(f"<b>{esc(g)}:</b> " + esc(" · ".join(items)), body))
    story.append(Paragraph("PROFESSIONAL EXPERIENCE", h))
    for tr in t.roles:
        role = roles[tr.role_id]
        left, right = _role_header(role)
        story.append(Paragraph(f"<b>{esc(left)}</b> &nbsp;&nbsp;<i>{esc(right)}</i>", body))
        story.append(ListFlowable([ListItem(Paragraph(esc(b.text), body), leftIndent=10) for b in tr.bullets],
                                  bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8))
    story.append(Paragraph("EDUCATION", h))
    for e in cv["education"]:
        story.append(Paragraph(f"<b>{esc(e['degree'])} — {esc(e['institution'])}</b> &nbsp;&nbsp;<i>{esc(e['dates'])} · {esc(e['location'])}</i>", body))
        if e.get("note"):
            story.append(Paragraph(esc(e["note"]), body))
    story.append(Paragraph("CREDENTIALS &amp; LANGUAGES", h))
    story.append(ListFlowable([ListItem(Paragraph(esc(c), body), leftIndent=10) for c in cv["credentials"]] +
                              [ListItem(Paragraph("Languages: " + esc(cv["languages"]), body), leftIndent=10)],
                              bulletType="bullet", start="•", leftIndent=12, bulletFontSize=8))
    path.parent.mkdir(parents=True, exist_ok=True)
    SimpleDocTemplate(str(path), pagesize=A4, leftMargin=16 * mm, rightMargin=16 * mm, topMargin=14 * mm,
                      bottomMargin=14 * mm, title=f"{cv['name']} - CV", author=cv["name"]).build(story)
    return path


# ------------------------------------------------------------------------------------------ Markdown (notes)
def render_notes_md(job, t: TailoredCV, engine: str, path: Path) -> Path:
    md = [f"# {job.title} — {job.company or 'n/a'}", f"- Source: {job.source} · {job.location} · posted {job.posted or 'n/a'}",
          f"- URL: {job.url}", f"- Rule score: {job.score} ({', '.join(job.score_reasons)})",
          f"- Fit (engine={engine}): {t.fit_score} — {t.fit_rationale}"]
    if job.red_flags:
        md.append("- Red flags: " + "; ".join(job.red_flags))
    md += ["", "## Keywords matched", ", ".join(t.keywords_matched) or "n/a", "", "## Gaps to address honestly",
           ", ".join(t.keywords_missing) or "none detected", "", f"## Email subject\n{t.email_subject}", "",
           "## Cover note", t.cover_note, "", "## LinkedIn note (<=300 chars)", t.linkedin_note]
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("\n".join(md), "utf-8")
    return path
